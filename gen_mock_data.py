import openpyxl, random, json, uuid
from datetime import date, timedelta

random.seed(42)

# 读取学生姓名
wb_src = openpyxl.load_workbook(r'C:\Users\DMJ\OneDrive\Desktop\七（4）班新生个人基本情况统计表.xlsx', data_only=True)
ws_src = wb_src['Sheet1']
students = []
for r in range(3, ws_src.max_row + 1):
    name = ws_src.cell(r, 2).value
    gender = ws_src.cell(r, 3).value
    if name:
        students.append({'name': str(name).strip(), 'gender': str(gender).strip() if gender else '男'})

# 模拟数据池
surnames = ['王','李','张','刘','陈','杨','赵','黄','周','吴','徐','孙','胡','朱','高','林','何','郭','马','罗','梁','宋','郑','谢','韩','唐','冯','于','董','萧','程','曹','袁','邓','许','傅','沈','曾','彭','吕']
male_names = ['伟','强','磊','军','洋','勇','杰','涛','明','超','平','刚','建华','文','辉','鹏','华','飞','鑫','波','斌','宇','浩','凯','健','俊','峰','阳','晨','博','豪','然','轩','睿','哲']
female_names = ['芳','娜','敏','静','丽','艳','娟','霞','秀英','慧','巧','美','英','梅','萍','红','玲','玉','燕','雪','琳','婷','欣','怡','佳','悦','璐','瑶','诗','琪','萱','涵','妍','彤','晴','曦','诺']
towns = ['叶邑镇','保安镇','辛店镇','龙泉乡','常村镇','夏李乡','田庄乡','龚店乡','邓李乡','水寨乡','廉村镇','洪庄杨乡','仙台镇','任店镇','马庄回族乡']
villages = ['小张庄','孟庄村','盆杨村','梅湾村','双庄村','老鸦张村','水郭村','菜庄村','蔡庄村','凤岭新村','杨岭庄村','思城村','段庄村','李庄村','王庄村','赵庄村','刘庄村','陈庄村','孙庄村','周庄村']
ethnicities = ['汉'] * 55 + ['满族', '回族', '蒙古族']

def random_phone():
    prefixes = ['135','136','137','138','139','150','151','152','158','159','182','183','186','187','188','176','177','199']
    return random.choice(prefixes) + ''.join([str(random.randint(0, 9)) for _ in range(8)])

def random_name(gender):
    s = random.choice(surnames)
    if gender == '男':
        g = random.choice(male_names)
    else:
        g = random.choice(female_names)
    if random.random() < 0.3:
        g += random.choice(male_names if gender == '男' else female_names)
    return s + g

def random_id_card(birth_date):
    area = '410422'
    birth = birth_date.strftime('%Y%m%d')
    seq = str(random.randint(100, 999))
    check = random.choice(['0','1','2','3','4','5','6','7','8','9','X'])
    return area + birth + seq + check

# 生成数据
app_students = []
for i, s in enumerate(students):
    start = date(2012, 9, 1)
    end = date(2013, 8, 31)
    birth = start + timedelta(days=random.randint(0, (end - start).days))
    age = date(2025, 9, 1).year - birth.year

    ethnicity = random.choice(ethnicities)
    father_name = random_name('男')
    father_phone = random_phone()
    mother_name = random_name('女')
    mother_phone = random_phone()
    address = '河南省-平顶山市-叶县-' + random.choice(towns) + random.choice(villages) + str(random.randint(1, 999)) + '号'
    id_card = random_id_card(birth)
    birth_str = birth.strftime('%Y年%m月')

    seat_row = i // 8 + 1
    seat_col = i % 8 + 1
    group_number = (i % 4) + 1

    # 写入 Excel
    row = i + 3
    ws_src.cell(row, 1, i + 1)
    ws_src.cell(row, 2, s['name'])
    ws_src.cell(row, 3, s['gender'])
    ws_src.cell(row, 4, ethnicity)
    ws_src.cell(row, 5, birth_str)
    ws_src.cell(row, 6, id_card)
    ws_src.cell(row, 7, age)
    ws_src.cell(row, 8, father_name)
    ws_src.cell(row, 9, father_phone)
    ws_src.cell(row, 10, mother_name)
    ws_src.cell(row, 11, mother_phone)
    ws_src.cell(row, 12, address)

    # App 学生数据（严格按照 Student 模型字段）
    notes = '父亲：' + father_name + ' ' + father_phone + '；母亲：' + mother_name + ' ' + mother_phone + '；民族：' + ethnicity + '；出生：' + birth_str + '；身份证：' + id_card
    app_students.append({
        'id': str(uuid.uuid4()),
        'name': s['name'],
        'studentNumber': '2025' + str(i + 1).zfill(3),
        'gender': s['gender'],
        'phone': '',
        'parentPhone': father_phone,
        'address': address,
        'groupNumber': group_number,
        'seatRow': seat_row,
        'seatCol': seat_col,
        'dormitory': '',
        'notes': notes
    })

# 保存 Excel
excel_path = r'D:\文件归档\桌面归档\02_我的项目\ClassTeacherApp\七（4）班学生信息_模拟数据.xlsx'
wb_src.save(excel_path)
print('Excel 已保存:', excel_path)

# 按照 AllDataBackup 真实结构生成 JSON
semester_id = str(uuid.uuid4())
backup = {
    'classInfo': {
        'className': '七（4）班',
        'grade': '七年级',
        'headTeacher': '',
        'subjects': ['语文', '数学', '英语', '政治', '历史', '地理', '生物', '体育', '音乐', '美术']
    },
    'students': app_students,
    'semesters': [
        {
            'id': semester_id,
            'name': '2025-2026学年第一学期',
            'shortName': '第1学期',
            'startDate': '2025-09-01T00:00:00Z',
            'endDate': '2026-01-31T00:00:00Z',
            'isCurrent': True
        }
    ],
    'exams': [],
    'scoreRecords': [],
    'courses': [],
    'dutyGroups': [],
    'todos': [],
    'notifications': [],
    'albumFolders': [],
    'albumPhotos': [],
    'photoFiles': {}
}
json_path = r'D:\文件归档\桌面归档\02_我的项目\ClassTeacherApp\七4班_模拟数据_备份.json'
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(backup, f, ensure_ascii=False, indent=2)
print('JSON 备份已保存:', json_path)
print('学生数:', len(app_students))
print('示例:', app_students[0]['name'], '-', app_students[0]['gender'], '-', app_students[0]['parentPhone'])
print('备注:', app_students[0]['notes'][:60])
